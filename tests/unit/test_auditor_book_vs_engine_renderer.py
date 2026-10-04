"""Tripwire: the auditor's book lens must read the ENGINE's real workbook clean.

The book lens reimplements the delivery contract's cell conventions rather
than importing the engine's renderer (independence). This test lives on the
engine side of the boundary and proves the two implementations agree: a
workbook rendered by the REAL ``render_workbook`` — bucket ordering,
subtotal furniture, the unprocessed section, status fills and all — must
produce zero findings for a faithful ledger, and a doctored cell must still
be caught. If the renderer's conventions ever change, this fails before any
nightly run cries wolf (or worse, goes quiet).
"""

from __future__ import annotations

from openpyxl import load_workbook

from auditor.evals.fixtures import add_invoice, make_context, make_ledger
from auditor.lenses import book
from core.agents.ap.workbook import render_workbook
from core.engine.config import WorkbookColumn

COLUMNS = [
    WorkbookColumn(label="Vendor", field="vendor"),
    WorkbookColumn(label="Invoice #", field="invoice_number"),
    WorkbookColumn(label="Amount", field="amount_cents"),
    WorkbookColumn(label="Status", field="status"),
    WorkbookColumn(label="Payment Date", field="payment_date"),
    WorkbookColumn(label="Check/Ref #", field="check_ref"),
    WorkbookColumn(label="QBO_Bill_ID", field=""),
]
AUDITOR_COLUMNS = [(c.label, c.field) for c in COLUMNS]


def _seed(conn):
    add_invoice(conn, vendor="Alpha", invoice_number="A-1", amount_cents=662400)
    add_invoice(
        conn,
        vendor="Beta",
        invoice_number="B-1",
        amount_cents=55217,
        status="Scheduled in bill pay",
        check_ref="4042",
    )
    add_invoice(
        conn,
        vendor="Gamma",
        invoice_number="G-1",
        amount_cents=874417,
        status="Paid",
        payment_date="2026-07-01",
    )
    add_invoice(conn, vendor="Delta", invoice_number="D-1", amount_cents=100, status="Cancelled")


def _render_with_engine(conn, path, *, unprocessed=True):
    rows = [dict(r) for r in conn.execute("SELECT * FROM ap_invoices ORDER BY id")]
    extra = (
        [{"file": "mystery-scan.pdf", "reason": "needs identification"}] if unprocessed else None
    )
    wb = render_workbook(rows, COLUMNS, legal_name="Test Co", unprocessed=extra)
    wb.save(str(path))


def _ctx(tmp_path, workbook_path):
    return make_context(
        tmp_path / "ledger",
        workbook_path=str(workbook_path),
        workbook_columns=AUDITOR_COLUMNS,
    )


def test_engine_rendered_workbook_reads_clean(tmp_path):
    conn = make_ledger(tmp_path / "ledger")
    _seed(conn)
    sheet = tmp_path / "book.xlsx"
    _render_with_engine(conn, sheet)
    ctx = _ctx(tmp_path, sheet)
    with ctx.ledger:
        findings = book.check(ctx)
    assert findings == [], [f.detail for f in findings]


def test_doctored_cell_in_engine_workbook_is_caught(tmp_path):
    conn = make_ledger(tmp_path / "ledger")
    _seed(conn)
    sheet = tmp_path / "book.xlsx"
    _render_with_engine(conn, sheet)
    # A hand edit on the delivered view: Beta's amount quietly shrinks.
    wb = load_workbook(sheet)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Beta":
            row[2].value = "$52.34"
    wb.save(str(sheet))
    ctx = _ctx(tmp_path, sheet)
    with ctx.ledger:
        findings = book.check(ctx)
    assert [f.condition for f in findings] == ["field-drift"]
    assert "Beta" in findings[0].subject
