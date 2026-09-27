"""Issue #166: the workbook view carries a generated Vendor Summary sheet.

The AR-side register used to hold a hand-frozen copy of this rollup, which
rotted (found stale at 2026-06-11 on 2026-09-01). A hand-copied view always
rots; the fix is a second sheet in the daily-regenerated workbook, which can
never be older than the AP Ledger sheet beside it.

Semantics: the summary rolls up exactly the rows the AP Ledger sheet shows,
per vendor spelling (1099/taxpayer grouping stays the auditor appendix's job).
Paid = "Paid"; Outstanding = payable statuses; Committed = committed statuses;
void/cancelled rows are excluded from counts and sums (no cash meaning).
Amounts are numeric cells with a currency format so a reader can sum them.
"""

from __future__ import annotations

from core.agents.ap.workbook import render_workbook
from core.engine.config import WorkbookColumn

_COLUMNS = [
    WorkbookColumn(label="Vendor", field="vendor"),
    WorkbookColumn(label="Invoice #", field="invoice_number"),
    WorkbookColumn(label="Amount", field="amount_cents"),
    WorkbookColumn(label="Status", field="status"),
]

_ROWS = [
    {"vendor": "Acme", "invoice_number": "1", "amount_cents": 10000, "status": "Paid"},
    {"vendor": "Acme", "invoice_number": "2", "amount_cents": 20000, "status": "Received"},
    {"vendor": "Acme", "invoice_number": "3", "amount_cents": 5000, "status": "Void - Duplicate"},
    {"vendor": "Beta", "invoice_number": "4", "amount_cents": 70000, "status": "Paid"},
    {"vendor": "Beta", "invoice_number": "5", "amount_cents": 40000, "status": "Scheduled"},
    {"vendor": "Gamma", "invoice_number": "6", "amount_cents": 1500, "status": "Cancelled"},
    {"vendor": "Delta", "invoice_number": "7", "amount_cents": 2500, "status": "Outstanding"},
]


def _summary(wb):
    return wb["Vendor Summary"]


def _data_rows(ws):
    """(vendor -> row values) for the data region, plus the TOTAL row values."""
    body: dict[str, tuple] = {}
    total = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "TOTAL":
            total = row
        elif row[0] is not None:
            body[row[0]] = row
    return body, total


def test_vendor_summary_sheet_exists_after_the_ledger_sheet():
    wb = render_workbook(_ROWS, _COLUMNS, legal_name="Acme Holdings LLC")
    assert wb.sheetnames == ["AP Ledger", "Vendor Summary"]
    assert wb.active.title == "AP Ledger"  # the daily view opens where it always did


def test_vendor_summary_headers_and_rollup_math():
    wb = render_workbook(_ROWS, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = _summary(wb)
    assert [c.value for c in ws[1]] == [
        "Vendor",
        "# Invoices",
        "Total Paid",
        "Outstanding (payable)",
        "Committed (scheduled)",
    ]
    body, _ = _data_rows(ws)
    # Acme: void row excluded from count and sums
    assert body["Acme"] == ("Acme", 2, 100.00, 200.00, 0.00)
    # Beta: paid + committed
    assert body["Beta"] == ("Beta", 2, 700.00, 0.00, 400.00)
    # Delta: Outstanding status is payable
    assert body["Delta"] == ("Delta", 1, 0.00, 25.00, 0.00)
    # Gamma had only a cancelled row: excluded entirely
    assert "Gamma" not in body


def test_vendor_summary_total_row_ties_and_is_bold():
    wb = render_workbook(_ROWS, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = _summary(wb)
    body, total = _data_rows(ws)
    assert total is not None
    assert total[1] == sum(v[1] for v in body.values())
    assert total[2] == sum(v[2] for v in body.values())
    assert total[3] == sum(v[3] for v in body.values())
    assert total[4] == sum(v[4] for v in body.values())
    total_row_idx = next(
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "TOTAL"
    )
    assert ws.cell(row=total_row_idx, column=1).font.bold is True
    assert total_row_idx == ws.max_row  # TOTAL is the last row


def test_vendor_summary_sorted_largest_first_deterministic():
    rows = _ROWS + [
        # same cash total as Delta: ties break alphabetically
        {"vendor": "Carol", "invoice_number": "8", "amount_cents": 2500, "status": "Paid"},
    ]
    wb = render_workbook(rows, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = _summary(wb)
    order = [
        ws.cell(row=r, column=1).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=1).value not in (None, "TOTAL")
    ]
    assert order == ["Beta", "Acme", "Carol", "Delta"]  # cash desc, then vendor asc


def test_vendor_summary_amounts_are_numeric_currency_cells():
    wb = render_workbook(_ROWS, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = _summary(wb)
    cell = ws.cell(row=2, column=3)  # first vendor's Total Paid
    assert isinstance(cell.value, float)  # numeric, so a reader can sum the column
    assert "#,##0.00" in cell.number_format
    assert isinstance(ws.cell(row=2, column=2).value, int)  # count stays an int


def test_vendor_summary_is_presentation_ready():
    from openpyxl.utils import get_column_letter

    wb = render_workbook(_ROWS, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = _summary(wb)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.cell(row=1, column=1).fill.fill_type == "solid"
    for idx in range(1, 6):
        assert ws.column_dimensions[get_column_letter(idx)].width > 0


def test_workbook_job_writes_both_sheets_through_guard(tmp_path):
    from openpyxl import load_workbook

    from core.agents.ap import store
    from core.engine.runner import resolve_ledger_root, run
    from core.ledger import Ledger

    ledger_dir = tmp_path / "data"
    out = tmp_path / "view" / "Engine_AP_Ledger.xlsx"
    root = resolve_ledger_root("demo", ledger_dir)
    with Ledger.open(root) as ledger:
        store.insert_invoice(
            ledger,
            tenant="demo",
            vendor="Acme",
            invoice_number="1",
            amount_cents=50000,
            status="Paid",
            shadow=True,
        )

    run(
        "demo",
        "ap",
        "workbook",
        shadow=False,
        params={"workbook_path": str(out)},
        ledger_dir=ledger_dir,
    )

    wb = load_workbook(str(out))
    assert wb.sheetnames == ["AP Ledger", "Vendor Summary"]
    ws = wb["Vendor Summary"]
    body, total = _data_rows(ws)
    assert body["Acme"] == ("Acme", 1, 500.00, 0.00, 0.00)
    assert total[2] == 500.00
