"""Half A of piece 2: the human-readable workbook view (delivery view).

Rows are colored by status group, the operator's at-a-glance unpaid / scheduled
/ paid: blank (payable), red (committed/scheduled), green (Paid/cleared), grey
(void). Columns come from tenant config so the sheet matches the legacy layout;
the output passes the write guard and carries the tenant's legal name
(invariant 10). 2026-06-26 cutover, Half A.
"""

from __future__ import annotations

from core.agents.ap.workbook import _cell_value, _status_fill, render_workbook
from core.engine.config import WorkbookColumn

_COLUMNS = [
    WorkbookColumn(label="Vendor", field="vendor"),
    WorkbookColumn(label="Invoice #", field="invoice_number"),
    WorkbookColumn(label="Amount", field="amount_cents"),
    WorkbookColumn(label="Status", field="status"),
    WorkbookColumn(label="Payment Date", field=""),  # no engine field -> blank
]


def test_status_fill_maps_groups_to_colors():
    assert _status_fill("Received") is None  # payable -> blank
    assert _status_fill("Outstanding") is None
    assert _status_fill("Scheduled") is not None  # committed -> red
    assert _status_fill("Scheduled") == _status_fill("Scheduled in bill pay")
    assert _status_fill("Paid") is not None  # settled-paid -> green
    assert _status_fill("Paid") != _status_fill("Scheduled")  # green != red
    void = _status_fill("Void - Duplicate")  # settled, never cleared -> grey
    assert void is not None and void not in (_status_fill("Paid"), _status_fill("Scheduled"))


def test_cell_value_formats_amount_and_blank_field():
    row = {"vendor": "Acme", "amount_cents": 110700, "status": "Received"}
    assert _cell_value(row, "amount_cents") == "$1,107.00"
    assert _cell_value(row, "") == ""  # a no-field column renders blank
    assert _cell_value(row, "vendor") == "Acme"
    assert _cell_value(row, "missing") == ""


def test_render_workbook_header_rows_colors_and_metadata():
    rows = [
        {"vendor": "Acme", "invoice_number": "1", "amount_cents": 50000, "status": "Scheduled"},
        {"vendor": "Beta", "invoice_number": "2", "amount_cents": 10000, "status": "Received"},
        {"vendor": "Gamma", "invoice_number": "3", "amount_cents": 20000, "status": "Paid"},
    ]
    wb = render_workbook(rows, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = wb.active

    assert [c.value for c in ws[1]] == ["Vendor", "Invoice #", "Amount", "Status", "Payment Date"]
    assert wb.properties.creator == "Acme Holdings LLC"  # invariant 10

    body = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(row=body["Acme"], column=1).fill.fill_type == "solid"  # scheduled -> filled
    assert ws.cell(row=body["Beta"], column=1).fill.fill_type in (None, "none")  # received -> blank


def test_workbook_has_cash_flow_subtotals_per_bucket():
    rows = [
        {"vendor": "A", "invoice_number": "1", "amount_cents": 10000, "status": "Received"},
        {"vendor": "B", "invoice_number": "2", "amount_cents": 20000, "status": "Outstanding"},
        {"vendor": "C", "invoice_number": "3", "amount_cents": 50000, "status": "Scheduled"},
        {"vendor": "D", "invoice_number": "4", "amount_cents": 70000, "status": "Paid"},
    ]
    wb = render_workbook(rows, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = wb.active
    # column 1 carries the bucket subtotal labels; column 3 carries the sums
    labels = {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=3).value
        for r in range(2, ws.max_row + 1)
    }
    assert labels.get("Pending total (2)") == "$300.00"  # 100 + 200
    assert labels.get("Scheduled total (1)") == "$500.00"
    assert labels.get("Paid total (1)") == "$700.00"
    # order is pending, then scheduled, then paid (actionable first)
    rownum = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    assert rownum["Pending total (2)"] < rownum["Scheduled total (1)"] < rownum["Paid total (1)"]


def test_workbook_shows_unprocessed_section():
    rows = [{"vendor": "A", "invoice_number": "1", "amount_cents": 10000, "status": "Received"}]
    unprocessed = [
        {"md5": "a" * 32, "file": "weird_scan.pdf", "reason": "extraction failed"},
        {"md5": "b" * 32, "file": "marketing_graphic_01.png", "reason": "image-only, needs OCR"},
    ]
    wb = render_workbook(rows, _COLUMNS, legal_name="Acme Holdings LLC", unprocessed=unprocessed)
    ws = wb.active
    col1 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "Unprocessed / needs identification (2)" in col1
    assert "weird_scan.pdf" in col1
    assert "marketing_graphic_01.png" in col1


def test_workbook_is_presentation_ready():
    from openpyxl.utils import get_column_letter

    rows = [
        {
            "vendor": "Acme Corporation With A Deliberately Long Name",
            "invoice_number": "1",
            "amount_cents": 50000,
            "status": "Received",
        }
    ]
    wb = render_workbook(rows, _COLUMNS, legal_name="Acme Holdings LLC")
    ws = wb.active

    assert ws.freeze_panes == "A2"  # header stays visible as the ledger grows
    assert ws.auto_filter.ref is not None  # sortable / filterable
    assert ws.column_dimensions[get_column_letter(1)].width > 0  # content-fit width, not default
    assert ws.cell(row=2, column=1).alignment.wrap_text is True  # long content wraps
    assert ws.cell(row=2, column=3).alignment.horizontal == "right"  # amount right-aligned
    assert ws.cell(row=1, column=1).font.bold is True  # header bold


def test_workbook_job_writes_view_through_guard(tmp_path):
    from openpyxl import load_workbook

    from core.agents.ap import store
    from core.engine.config import load_tenant
    from core.engine.runner import resolve_ledger_root, run
    from core.ledger import Ledger

    ledger_dir = tmp_path / "data"
    out = tmp_path / "view" / "Engine_AP_Ledger.xlsx"
    root = resolve_ledger_root("demo", ledger_dir)
    with Ledger.open(root) as ledger:
        store.insert_invoice(
            ledger,
            tenant="demo",
            vendor="Acme",
            invoice_number="1",
            amount_cents=50000,
            status="Scheduled",
            shadow=True,
        )

    run(
        "demo",
        "ap",
        "workbook",
        shadow=False,
        params={"workbook_path": str(out)},
        ledger_dir=ledger_dir,
    )

    assert out.exists()
    ws = load_workbook(str(out)).active
    # The header IS the tenant's [ap].workbook_columns labels, read from the
    # config rather than spelled out: since row 7.19 the demo is rendered from
    # the archetype template, whose unit column carries the archetype's own
    # vocabulary ("Project" for A, "Job" for B).
    labels = [c.label for c in load_tenant("demo").ap.workbook_columns]
    assert [c.value for c in ws[1]] == labels
    assert labels[:3] == ["Vendor", "Invoice #", "Amount"] and labels[-1] == "Status"
    assert ws.cell(row=2, column=1).value == "Acme"
    assert ws.cell(row=2, column=3).value == "$500.00"
