"""Shadow parity diff: engine shadow rows vs the legacy AP ledger workbook.

Read-only toward the legacy workbook (it is a protected surface; this module
never writes it). The report is designed to be judged in under two minutes:
verdict and counts first, then only the rows that need eyes.

Comparison semantics:
- Join key: canonical vendor + invoice number (identity, never amount).
- The legacy workbook holds the whole year; the engine ledger holds what
  shadow intake has processed. Legacy-only rows are therefore restricted to
  invoices dated on/after the shadow window's start, so old history does not
  drown the signal.
- Statuses compare by GROUP (payable / committed / settled), because the
  engine starts every row at Received while a human may flip the legacy row
  the same day. A group mismatch is a real finding; a same-group wording
  difference is not.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from ...engine.guard import WriteGuard
from ...ledger import Ledger
from .registry import VendorRegistry, canonical_vendor, load_vendor_registry
from .status import COMMITTED_STATUSES, SETTLED_STATUSES

LEDGER_SHEET = "AP Ledger"
_DIVIDER_TOKENS = ("OUTSTANDING", "SCHEDULED", "PAID / VOID")

# Vendor identity (canonical name + ledger aliases, trailing parenthetical
# dropped) is computed by ``registry.canonical_vendor`` so the parity diff and
# the three-way payment verification share one matcher and cannot drift apart
# (2026-07-01). Both the 2026-06-16 parenthetical finding and the 2026-06-23
# ledger-alias finding live in that function now.


def _status_group(status: str) -> str:
    if status in COMMITTED_STATUSES:
        return "committed"
    if status in SETTLED_STATUSES:
        return "settled"
    return "payable"


def _cents(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int((Decimal(str(value)) * 100).to_integral_value())


@dataclass
class LegacyRow:
    vendor: str
    invoice_number: str
    amount_cents: int | None
    status: str
    invoice_date: str = ""

    @property
    def key(self) -> tuple[str, str]:
        return (canonical_vendor(self.vendor), self.invoice_number.strip().lower())


@dataclass
class ParityReport:
    matched: int = 0
    mismatches: list[str] = field(default_factory=list)
    engine_only: list[str] = field(default_factory=list)
    legacy_only: list[str] = field(default_factory=list)
    # Engine-only and legacy-only rows that share an exact invoice number are
    # almost certainly one vendor under two spellings. Surfaced as a
    # confirm-to-link suggestion, never auto-merged (2026-06-23 shadow finding).
    alias_candidates: list[str] = field(default_factory=list)
    # Legacy rows with no invoice date cannot be windowed (recurring
    # auto-debits, card paydowns, time-based rows). They are listed for a
    # one-time review but never block the daily verdict; the engine does not
    # source these row types from the landing folder at all.
    undated_legacy: list[str] = field(default_factory=list)

    @property
    def is_parity(self) -> bool:
        return not (self.mismatches or self.engine_only or self.legacy_only)


def read_legacy_rows(xlsx_path: str | Path) -> list[LegacyRow]:
    """Read the legacy workbook's AP Ledger sheet, read-only, skipping bands."""
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    try:
        ws = wb[LEDGER_SHEET]
        rows: list[LegacyRow] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            vendor = str(row[0] or "").strip()
            status = str(row[8] or "").strip()
            if not vendor or not status:
                continue  # divider bands and blanks carry no status
            if any(token in vendor.upper() for token in _DIVIDER_TOKENS):
                continue
            invoice_date = row[2]
            if hasattr(invoice_date, "date"):
                invoice_date = invoice_date.date().isoformat()
            rows.append(
                LegacyRow(
                    vendor=vendor,
                    invoice_number=str(row[1] or "").strip(),
                    amount_cents=_cents(row[4]),
                    status=status,
                    invoice_date=str(invoice_date or ""),
                )
            )
        return rows
    finally:
        wb.close()


def engine_shadow_rows(ledger: Ledger, tenant: str) -> list[dict]:
    rows = ledger.conn.execute(
        "SELECT vendor, invoice_number, amount_cents, status, invoice_date "
        "FROM ap_invoices WHERE tenant = ? AND shadow = 1 ORDER BY id",
        (tenant,),
    ).fetchall()
    return [dict(r) for r in rows]


def _join_key(vendor: str, invoice_number: str, registry: VendorRegistry | None) -> tuple[str, str]:
    return (canonical_vendor(vendor, registry), invoice_number.strip().lower())


def diff(
    engine_rows: list[dict],
    legacy_rows: list[LegacyRow],
    *,
    since: str = "",
    registry: VendorRegistry | None = None,
) -> ParityReport:
    report = ParityReport()
    legacy_by_key = {_join_key(r.vendor, r.invoice_number, registry): r for r in legacy_rows}

    engine_unmatched: list[dict] = []
    for row in engine_rows:
        key = _join_key(row["vendor"], row["invoice_number"], registry)
        legacy = legacy_by_key.pop(key, None)
        if legacy is None:
            engine_unmatched.append(row)
            report.engine_only.append(
                f"{row['vendor']} / {row['invoice_number']} (engine has it, legacy ledger does not)"
            )
            continue
        problems = []
        if legacy.amount_cents is not None and legacy.amount_cents != row["amount_cents"]:
            problems.append(
                f"amount: engine {row['amount_cents'] / 100:.2f} vs "
                f"legacy {legacy.amount_cents / 100:.2f}"
            )
        if _status_group(row["status"]) != _status_group(legacy.status):
            problems.append(f"status group: engine {row['status']!r} vs legacy {legacy.status!r}")
        if problems:
            report.mismatches.append(
                f"{row['vendor']} / {row['invoice_number']}: " + "; ".join(problems)
            )
        else:
            report.matched += 1

    legacy_unmatched: list[LegacyRow] = []
    for legacy in legacy_by_key.values():
        if since and not legacy.invoice_date:
            report.undated_legacy.append(
                f"{legacy.vendor} / {legacy.invoice_number or '(no ref)'} ({legacy.status})"
            )
            continue
        if since and legacy.invoice_date < since:
            continue  # pre-window history is not a parity miss
        legacy_unmatched.append(legacy)
        report.legacy_only.append(
            f"{legacy.vendor} / {legacy.invoice_number} "
            f"(invoice date {legacy.invoice_date or 'unknown'}; legacy has it, engine missed it)"
        )

    _detect_alias_candidates(report, engine_unmatched, legacy_unmatched, registry)
    return report


def _detect_alias_candidates(
    report: ParityReport,
    engine_unmatched: list[dict],
    legacy_unmatched: list[LegacyRow],
    registry: VendorRegistry | None,
) -> None:
    """Surface engine-only/legacy-only pairs that share an exact invoice number.

    Two vendors do not issue the same invoice number, so a shared number is a
    high-confidence signal that two spellings are one vendor. This is the only
    matching done; the suggestion is for a human to confirm, after which the
    legacy spelling is added to that vendor's ``ledger_aliases``. Nothing is
    auto-merged.
    """
    legacy_by_invoice: dict[str, list[LegacyRow]] = {}
    for legacy in legacy_unmatched:
        legacy_by_invoice.setdefault(legacy.invoice_number.strip().lower(), []).append(legacy)

    seen: set[tuple[str, str, str]] = set()
    for row in engine_unmatched:
        invoice = row["invoice_number"].strip().lower()
        if not invoice:
            continue
        for legacy in legacy_by_invoice.get(invoice, []):
            if canonical_vendor(row["vendor"], registry) == canonical_vendor(
                legacy.vendor, registry
            ):
                continue  # already one identity; not a candidate
            pair = (legacy.vendor, row["vendor"], invoice)
            if pair in seen:
                continue
            seen.add(pair)
            report.alias_candidates.append(
                f"'{legacy.vendor}' (legacy) = '{row['vendor']}' (engine), "
                f"seen on invoice {row['invoice_number']}"
            )


def render_markdown(report: ParityReport, *, tenant: str, since: str) -> str:
    verdict = "PARITY" if report.is_parity else "NEEDS ATTENTION"
    lines = [
        f"# Shadow parity report: {tenant} / ap",
        "",
        f"Generated: {datetime.now(UTC).isoformat(timespec='seconds')}",
        f"Window: invoices dated on/after {since or '(all)'}",
        "",
        f"## Verdict: {verdict}",
        "",
        f"- matched rows: {report.matched}",
        f"- field mismatches: {len(report.mismatches)}",
        f"- engine-only rows: {len(report.engine_only)}",
        f"- legacy-only rows (in window): {len(report.legacy_only)}",
        f"- candidate aliases (confirm to link): {len(report.alias_candidates)}",
    ]
    for title, items in (
        ("Field mismatches", report.mismatches),
        ("Engine-only", report.engine_only),
        ("Legacy-only (in window)", report.legacy_only),
    ):
        if items:
            lines += ["", f"## {title}", ""]
            lines += [f"- {item}" for item in items]
    if report.alias_candidates:
        lines += [
            "",
            "## Candidate vendor aliases (confirm to link)",
            "",
            "These engine-only and legacy-only rows share an exact invoice number,",
            "so the two names are almost certainly one vendor under two spellings.",
            "Confirm each, then add the legacy spelling to that vendor's",
            "ledger_aliases in vendors.toml and the pair reconciles next run.",
            "",
        ]
        lines += [f"- {item}" for item in report.alias_candidates]
    if report.undated_legacy:
        lines += [
            "",
            "## Undated legacy rows (informational, not counted in the verdict)",
            "",
            "These ledger rows carry no invoice date, so they cannot be windowed.",
            "They are row types the engine does not source from the landing folder",
            "(recurring auto-debits, card paydowns, time-based rows). Eyeball once",
            "at the start of the shadow period; after that they are noise.",
            "",
        ]
        lines += [f"- {item}" for item in report.undated_legacy]
    if report.is_parity:
        lines += ["", "Nothing needs your eyes today."]
    lines.append("")
    return "\n".join(lines)


def _tenant_registry(tenant_slug: str) -> VendorRegistry:
    """Load the tenant's vendor registry; empty registry if it has none."""
    try:
        from ...engine.config import tenant_dir

        path = tenant_dir(tenant_slug) / "vendors.toml"
        if path.exists():
            return load_vendor_registry(path)
    except (FileNotFoundError, OSError, ValueError):
        pass
    return VendorRegistry()


def run_shadow_diff(
    *,
    tenant_slug: str,
    ledger: Ledger,
    legacy_xlsx: str | Path,
    guard: WriteGuard,
    since: str = "",
    out_path: str | Path | None = None,
    registry: VendorRegistry | None = None,
) -> tuple[ParityReport, Path | None]:
    """Produce the parity report; optionally write it inside the ledger root."""
    legacy = read_legacy_rows(legacy_xlsx)
    engine = engine_shadow_rows(ledger, tenant_slug)
    if registry is None:
        registry = _tenant_registry(tenant_slug)
    report = diff(engine, legacy, since=since, registry=registry)
    written: Path | None = None
    if out_path is not None:
        target = guard.check_write(out_path)  # never a protected surface
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            render_markdown(report, tenant=tenant_slug, since=since), encoding="utf-8"
        )
        written = target
    return report, written
