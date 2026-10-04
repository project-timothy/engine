"""#326 end to end: an AP write anywhere re-renders the delivered workbook once.

The 2026-09-17 window: six AP rows written after the morning render (two of
them by hand, outside any job run) left the delivered sheet wrong about
$91,700 for twenty hours and raised six CRITICALs that fixed themselves the
next morning. The store now marks the view stale after every AP write; the end
of any engine invocation (and the 15-minute retries job) renders it once. A
failed render never touches the AP write and never changes an exit code.
"""

from __future__ import annotations

from openpyxl import load_workbook

from core.agents.ap import store, view_refresh
from core.engine.cli import main
from core.engine.runner import resolve_ledger_root, run
from core.ledger import Ledger


def _book(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # the demo's workbook_path is relative to the run dir
    ledger_dir = tmp_path / "data"
    run("demo", "ap", "workbook", ledger_dir=ledger_dir)  # the morning render
    return (
        ledger_dir,
        resolve_ledger_root("demo", ledger_dir),
        tmp_path / "demo-data" / "reports" / "ap-ledger.xlsx",
    )


def _numbers(xlsx) -> list[str]:
    wb = load_workbook(xlsx, read_only=True)
    cells = [str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row]
    return [c for c in cells if c.startswith("HB-")]


def _hand_backfill(root, number: str) -> None:
    """The 2026-09-17 path: store.insert_invoice outside any job run."""
    with Ledger.open(root) as ledger:
        store.insert_invoice(
            ledger,
            tenant="demo",
            vendor="Acme Tooling",
            invoice_number=number,
            amount_cents=240000,
            status="Paid",
        )


def test_a_hand_backfill_marks_the_view_stale(tmp_path, monkeypatch):
    ledger_dir, root, _ = _book(tmp_path, monkeypatch)
    assert not view_refresh.is_stale(root)
    _hand_backfill(root, "HB-1")
    assert view_refresh.is_stale(root)


def test_the_next_engine_invocation_renders_it_once(tmp_path, monkeypatch, capsys):
    ledger_dir, root, xlsx = _book(tmp_path, monkeypatch)
    _hand_backfill(root, "HB-1")
    _hand_backfill(root, "HB-2")
    assert _numbers(xlsx) == []
    code = main(["run", "demo", "demo", "ingest", "--ledger-dir", str(ledger_dir)])
    assert code == 0
    assert sorted(_numbers(xlsx)) == ["HB-1", "HB-2"]
    assert not view_refresh.is_stale(root)
    assert "view refreshed" in capsys.readouterr().err


def test_the_retries_job_catches_it_within_fifteen_minutes(tmp_path, monkeypatch):
    ledger_dir, root, xlsx = _book(tmp_path, monkeypatch)
    _hand_backfill(root, "HB-1")
    assert main(["jobs", "resume", "demo", "--ledger-dir", str(ledger_dir)]) == 0
    assert _numbers(xlsx) == ["HB-1"]


def test_a_shadow_invocation_renders_nothing(tmp_path, monkeypatch):
    ledger_dir, root, xlsx = _book(tmp_path, monkeypatch)
    _hand_backfill(root, "HB-1")
    main(["run", "demo", "demo", "ingest", "--shadow", "--ledger-dir", str(ledger_dir)])
    assert _numbers(xlsx) == [] and view_refresh.is_stale(root)


def test_a_shadow_workbook_run_writes_no_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    out = run("demo", "ap", "workbook", ledger_dir=tmp_path / "data", shadow=True)
    assert out.status == "ok"
    assert not (tmp_path / "demo-data" / "reports" / "ap-ledger.xlsx").exists()


def test_a_failed_render_leaves_the_write_and_the_exit_code(tmp_path, monkeypatch, capsys):
    ledger_dir, root, xlsx = _book(tmp_path, monkeypatch)
    _hand_backfill(root, "HB-1")

    def refused(tenant, *, ledger_dir=None):
        raise RuntimeError("openpyxl refused the workbook write")

    monkeypatch.setattr(view_refresh, "_render", refused)
    code = main(["run", "demo", "demo", "ingest", "--ledger-dir", str(ledger_dir)])
    assert code == 0
    assert "openpyxl refused the workbook write" in capsys.readouterr().err
    with Ledger.open(root) as ledger:
        assert store.find_invoice(ledger, "demo", "Acme Tooling", "HB-1") is not None
    assert view_refresh.is_stale(root)  # tried again at the next invocation


def test_the_morning_render_clears_the_mark(tmp_path, monkeypatch):
    ledger_dir, root, xlsx = _book(tmp_path, monkeypatch)
    _hand_backfill(root, "HB-1")
    run("demo", "ap", "workbook", ledger_dir=ledger_dir)
    assert not view_refresh.is_stale(root)
    assert _numbers(xlsx) == ["HB-1"]
