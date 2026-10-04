"""Lens 4 — book integrity: the delivered workbook tells the ledger's truth.

The workbook is a delivery view, a photograph of the ledger (engine
invariant 1). Any drift between them means the view LIES about the book —
the exact failure that invariant forbids, and the highest-value single
check the auditor runs. Every derivation here is the auditor's own:

- expected rows recomputed from ``ap_invoices`` with its own SQL;
- the delivered sheet opened with its own openpyxl reader;
- cell conventions (the dollar format, blank-for-empty) reimplemented from
  the delivery contract, never imported. The engine-side integration test
  in ``tests/unit/test_auditor_book_vs_engine_renderer.py`` renders a real
  workbook with the ENGINE's code and asserts this lens reads it clean, so
  the two implementations cannot drift silently.

Rows match on vendor + invoice number (amount as a tiebreak for repeats);
compared fields are the design's four: amount (integer cents, never
float), status, payment date, check ref. Subtotal lines, the unprocessed
section, and the header are the view's own furniture, not rows.
"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path

from ..findings import Finding
from . import AuditContext

LENS = "book"

# The auditor's own copy of the status vocabulary (see status_coherence).
KNOWN_STATUSES = frozenset(
    {
        "Received",
        "Approved",
        "Outstanding",
        "Scheduled",
        "Scheduled in bill pay",
        "Paid",
        "Void - Already Paid",
        "Void - Duplicate",
        "Cancelled",
    }
)

COMPARED_FIELDS = ("amount_cents", "status", "payment_date", "check_ref")


def dollars(cents: int) -> str:
    """The delivery contract's amount rendering, reimplemented."""
    return f"${cents / 100:,.2f}"


def parse_dollars(text: str) -> int | None:
    """A sheet amount back to integer cents; None when it is not an amount."""
    cleaned = str(text).replace("$", "").replace(",", "").strip()
    if not cleaned:
        return None
    try:
        return int((Decimal(cleaned) * 100).to_integral_value())
    except InvalidOperation:
        return None


def _expected_rows(ctx: AuditContext) -> list[dict]:
    return ctx.ledger.query(
        "SELECT id, vendor, invoice_number, amount_cents, status, "
        "COALESCE(payment_date, '') AS payment_date, check_ref "
        "FROM ap_invoices WHERE tenant = ? ORDER BY id",
        (ctx.tenant.slug,),
    )


def _sheet_rows(path: Path, columns: list[tuple[str, str]]) -> list[dict]:
    """Data rows from the delivered sheet, keyed by engine field name."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[dict] = []
        fields = [field for _, field in columns]
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            values = ["" if v is None else str(v) for v in excel_row]
            values += [""] * (len(fields) - len(values))
            row = {field: value for field, value in zip(fields, values, strict=False) if field}
            rows.append(row)
        return rows
    finally:
        wb.close()


def _is_data_row(row: dict, expected_vendors: set[str]) -> bool:
    """The view's furniture (subtotals, dividers, the unprocessed section)
    carries no invoice number; a data row does — or is a known vendor's row
    in a known status (the empty-invoice-number edge)."""
    if row.get("invoice_number", ""):
        return True
    return row.get("vendor", "") in expected_vendors and row.get("status", "") in KNOWN_STATUSES


def _key(vendor: str, invoice_number: str) -> tuple[str, str]:
    return (vendor.strip(), invoice_number.strip())


def check(ctx: AuditContext) -> list[Finding]:
    if not ctx.tenant.workbook_path or not ctx.tenant.workbook_columns:
        return []  # tenant delivers no workbook view
    tables = ctx.ledger.query(
        "SELECT name FROM sqlite_master WHERE type='table' AND name='ap_invoices'"
    )
    if not tables:
        return []  # tenant has no AP book to compare

    expected = _expected_rows(ctx)
    path = Path(ctx.tenant.workbook_path).expanduser()
    if not path.exists():
        if not expected:
            return []
        return [
            Finding(
                lens=LENS,
                subject="workbook",
                condition="missing",
                severity="CRITICAL",
                detail=f"the ledger holds {len(expected)} row(s) but the delivered "
                f"workbook does not exist at {path}",
            )
        ]

    sheet = _sheet_rows(path, ctx.tenant.workbook_columns)
    expected_vendors = {str(r["vendor"]) for r in expected}
    data_rows = [r for r in sheet if _is_data_row(r, expected_vendors)]

    findings: list[Finding] = []
    remaining: dict[tuple[str, str], list[dict]] = {}
    for row in data_rows:
        remaining.setdefault(_key(row.get("vendor", ""), row.get("invoice_number", "")), []).append(
            row
        )

    for row in expected:
        subject = f"invoice #{row['id']} {row['vendor']} / {row['invoice_number']}"
        candidates = remaining.get(_key(str(row["vendor"]), str(row["invoice_number"])), [])
        if not candidates:
            findings.append(
                Finding(
                    lens=LENS,
                    subject=subject,
                    condition="missing-from-sheet",
                    severity="CRITICAL",
                    detail=f"the ledger holds this row ({dollars(row['amount_cents'])}, "
                    f"{row['status']}) but the delivered sheet does not show it",
                )
            )
            continue
        # Prefer the candidate with the matching amount (repeats tiebreak).
        match = next(
            (
                c
                for c in candidates
                if parse_dollars(c.get("amount_cents", "")) == row["amount_cents"]
            ),
            candidates[0],
        )
        candidates.remove(match)
        drifts: list[str] = []
        sheet_cents = parse_dollars(match.get("amount_cents", ""))
        if sheet_cents != row["amount_cents"]:
            drifts.append(
                f"amount: ledger {dollars(row['amount_cents'])} vs sheet "
                f"{match.get('amount_cents', '') or '(blank)'}"
            )
        if match.get("status", "") != str(row["status"]):
            drifts.append(f"status: ledger {row['status']!r} vs sheet {match.get('status', '')!r}")
        if match.get("payment_date", "") != str(row["payment_date"]):
            drifts.append(
                f"payment date: ledger {row['payment_date'] or '(none)'} vs sheet "
                f"{match.get('payment_date', '') or '(blank)'}"
            )
        if match.get("check_ref", "") != str(row["check_ref"]):
            drifts.append(
                f"check ref: ledger {row['check_ref'] or '(none)'} vs sheet "
                f"{match.get('check_ref', '') or '(blank)'}"
            )
        if drifts:
            findings.append(
                Finding(
                    lens=LENS,
                    subject=subject,
                    condition="field-drift",
                    severity="CRITICAL",
                    detail="the delivered sheet disagrees with the ledger — " + "; ".join(drifts),
                )
            )

    for key, leftovers in remaining.items():
        for row in leftovers:
            vendor, number = key
            findings.append(
                Finding(
                    lens=LENS,
                    subject=f"sheet row {vendor} / {number}",
                    condition="not-in-ledger",
                    severity="CRITICAL",
                    detail=f"the delivered sheet shows this row "
                    f"({row.get('amount_cents', '') or 'no amount'}, "
                    f"{row.get('status', '') or 'no status'}) but the ledger does not hold it; "
                    "the view was edited by hand or generated from something else",
                )
            )
    return findings
