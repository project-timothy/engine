"""Lens 11 — PO watch: every archived customer PO reaches its two homes.

The mail fetch already captures every purchasing-portal PO email attachment
into the landing archive as ``PO-*.pdf`` — capture is covered. What can go
stale unnoticed for months is the ROUTING: the copy into the canonical PO folder and the row
in the receivables register's Open-POs sheet. This lens recomputes both
joins nightly from the filenames alone (no PDF parsing): a PO number seen
in the archive must appear in some filename in the PO home and in the
register sheet's first column.

Read-only, findings only — filing the PDF and appending the row stay
human/session judgments (issue #111's non-goal: no automated writes).
"""

from __future__ import annotations

import re
from pathlib import Path

from ..findings import Finding
from . import AuditContext
from .filing import tree_missing

LENS = "po-watch"

# PO-<series-number>_v<ver>_<YYYYMMDD>.pdf — the canonical shape.
_PO_NAME = re.compile(r"^PO-(?P<num>.+?)_v\d+.*\.(pdf|PDF)$")


def _po_number(name: str) -> str | None:
    m = _PO_NAME.match(name)
    return m.group("num") if m else None


def _register_numbers(xlsx: Path, sheet: str) -> set[str] | None:
    """First-column values of the register sheet, or None if unreadable."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(xlsx), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if sheet not in wb.sheetnames:
            return None
        numbers: set[str] = set()
        for row in wb[sheet].iter_rows(min_col=1, max_col=1, values_only=True):
            value = row[0]
            if value is not None:
                numbers.add(str(value).strip())
        return numbers
    finally:
        wb.close()


def check(ctx: AuditContext) -> list[Finding]:
    if not ctx.tenant.po_watch_enabled:
        return []
    archive = Path(ctx.tenant.po_archive_dir).expanduser()
    home = Path(ctx.tenant.po_home_dir).expanduser()
    register = Path(ctx.tenant.po_register_xlsx).expanduser()
    sheet = ctx.tenant.po_register_sheet
    if not archive.is_dir():
        # A configured archive that is not there is not "nothing captured
        # yet" (04-F4): say so and stop, nothing else is checkable.
        return [tree_missing(LENS, "PO archive", archive)]

    archived: dict[str, str] = {}  # number -> archive filename (first seen)
    for path in sorted(archive.glob("*/PO-*.pdf")) + sorted(archive.glob("PO-*.pdf")):
        num = _po_number(path.name)
        if num and num not in archived:
            archived[num] = path.name
    if not archived:
        return []

    # The home is the PO root with one folder per customer (issue #158):
    # scan the root and each customer subfolder, one level deep, so a new
    # customer's folder is covered with zero config. Underscore
    # folders (_inbound, _quotes) are staging, not filing — a PO sitting there
    # has not reached its home and must still warn.
    findings: list[Finding] = []
    home_numbers: set[str] = set()
    home_checkable = home.is_dir()
    if not home_checkable:
        # Name the missing tree once and leave the home half unchecked, the
        # register-unreadable shape; a per-PO "has no copy" would claim a
        # fact about a folder never read (04-F4).
        findings.append(tree_missing(LENS, "PO home", home))
    else:
        candidates = list(home.glob("PO-*.pdf"))
        for child in sorted(home.iterdir()):
            if child.is_dir() and not child.name.startswith("_"):
                candidates.extend(child.glob("PO-*.pdf"))
        for path in candidates:
            num = _po_number(path.name)
            if num:
                home_numbers.add(num)

    register_numbers = _register_numbers(register, sheet)

    for num, filename in sorted(archived.items()):
        if home_checkable and num not in home_numbers:
            findings.append(
                Finding(
                    lens=LENS,
                    subject=num,
                    condition="not-in-po-home",
                    severity="WARN",
                    detail=f"{filename} is archived but {home} has no copy; "
                    "file it to the canonical PO folder",
                )
            )
        if register_numbers is not None and num not in register_numbers:
            findings.append(
                Finding(
                    lens=LENS,
                    subject=num,
                    condition="not-in-register",
                    severity="WARN",
                    detail=f"PO {num} is missing from the register's "
                    f"'{sheet}' sheet; append the row (attribution stays a "
                    "human judgment — [NEEDS REVIEW] in the PN column if unsure)",
                )
            )
    if register_numbers is None:
        findings.append(
            Finding(
                lens=LENS,
                subject=register.name or "(register unset)",
                condition="register-unreadable",
                severity="WARN",
                detail=f"could not read sheet '{sheet}' from {register}; "
                "the register half of the PO join was not checked tonight",
            )
        )
    return findings
