"""The close packet: the month's evidence, one workbook, regenerated on
every render and never hand-edited.

Sheets: P&L for the month, balance sheet at month-end, AP aging, AR aging,
the bank tie-out, payroll entries, owner transactions, and the checklist
with a sign-off line. Numbers come from the accounting system's reports and
the ledger; report sheets render the report's own row tree (indented by
section) so the packet reads like the statements the owner already knows.
Author/company metadata is the tenant's legal name (invariant 10);
openpyxl per the house rule; the write passes the guard.
"""

from __future__ import annotations

from pathlib import Path

from ...engine.guard import WriteGuard
from .checks import (
    CloseContext,
    _cents_str,
    _committed_not_cleared,
    _register_balance_cents,
    month_bounds,
)
from .schema import CheckResult

_HEADER_FILL = "FFD9E1F2"
_STATUS_FILLS = {"OK": "FFC6EFCE", "WARN": "FFFFE699", "BLOCK": "FFFFC7CE"}


def _style_header(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)


def _fit_columns(ws, cap: int = 60) -> None:
    from openpyxl.utils import get_column_letter

    for idx in range(1, ws.max_column + 1):
        longest = max(
            (len(str(ws.cell(row=r, column=idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[get_column_letter(idx)].width = max(10, min(longest + 2, cap))


def _report_sheet(wb, title: str, report: dict) -> None:
    """A QBO report's row tree, section-indented, as the owner would read it."""
    from ...adapters.qbo import report_rows

    ws = wb.create_sheet(title)
    columns = [
        str(c.get("ColTitle", "") or "") for c in (report.get("Columns") or {}).get("Column", [])
    ]
    ws.append(columns or ["Line", "Amount"])
    for row in report_rows(report):
        cells = list(row["cells"])
        if cells:
            cells[0] = ("    " * len(row["path"])) + cells[0]
        ws.append(cells)
    _style_header(ws)
    _fit_columns(ws)


def _tie_out_sheet(wb, ctx: CloseContext) -> None:
    ws = wb.create_sheet("Bank Tie-Out")
    ws.append(["Item", "Amount"])
    register = _register_balance_cents(ctx)
    statement = ctx.statement_balance_cents
    ws.append(
        [
            "Statement ending balance",
            _cents_str(statement) if statement is not None else "(not provided)",
        ]
    )
    ws.append(
        [
            "Register balance at month-end",
            _cents_str(register) if register is not None else "(not found)",
        ]
    )
    if statement is not None and register is not None:
        ws.append(["Delta", _cents_str(statement - register)])
    ws.append([])
    ws.append(["Committed, not yet cleared (context: promised money, visible nowhere in QBO)"])
    for line in _committed_not_cleared(ctx):
        ws.append([line])
    _style_header(ws)
    _fit_columns(ws, cap=90)


def _payroll_sheet(wb, ctx: CloseContext) -> None:
    ws = wb.create_sheet("Payroll")
    ws.append(["Date", "Doc", "Note"])
    start, end = month_bounds(ctx.month)
    for je in ctx.qbo.fetch_journal_entries(start=start, end=end):
        ws.append(
            [
                str(je.get("TxnDate", "")),
                str(je.get("DocNumber", "")),
                str(je.get("PrivateNote", "")),
            ]
        )
    _style_header(ws)
    _fit_columns(ws)


def _owner_sheet(wb, ctx: CloseContext) -> None:
    ws = wb.create_sheet("Owner Transactions")
    ws.append(["Date", "Payee", "Amount", "Coding"])
    start, end = month_bounds(ctx.month)
    owners = [n.lower() for n in (ctx.owner_names or [])]
    for txn in ctx.qbo.fetch_purchases(start=start, end=end):
        payee = str(((txn.get("EntityRef") or {}).get("name", "")) or "")
        if payee and not any(o in payee.lower() for o in owners):
            continue
        accounts = sorted(
            {
                str(
                    (line.get("AccountBasedExpenseLineDetail") or {})
                    .get("AccountRef", {})
                    .get("name", "")
                )
                for line in txn.get("Line", [])
            }
            - {""}
        )
        ws.append(
            [
                str(txn.get("TxnDate", "")),
                payee or "(no payee)",
                str(txn.get("TotalAmt", "")),
                ", ".join(accounts) or "(no account)",
            ]
        )
    _style_header(ws)
    _fit_columns(ws)


def _checklist_sheet(wb, results: list[CheckResult], month: str) -> None:
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet("Checklist + Sign-off")
    ws.append(["Check", "Status", "Evidence"])
    for result in results:
        ws.append([result.name, result.status, result.summary])
        fill = _STATUS_FILLS.get(result.status)
        if fill:
            ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=fill)
    ws.append([])
    ws.append([f"Close of {month} reviewed and approved."])
    ws.append(["Owner sign-off:", "", "Date:"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    _style_header(ws)
    _fit_columns(ws, cap=100)


def render_packet(
    ctx: CloseContext,
    results: list[CheckResult],
    *,
    legal_name: str,
    out_path: str | Path,
    guard: WriteGuard,
) -> Path:
    from openpyxl import Workbook

    start, end = month_bounds(ctx.month)
    wb = Workbook()
    wb.remove(wb.active)  # sheets are created in packet order below

    _checklist_sheet(wb, results, ctx.month)
    _report_sheet(
        wb, "P&L", ctx.qbo.fetch_report("ProfitAndLoss", {"start_date": start, "end_date": end})
    )
    _report_sheet(
        wb,
        "Balance Sheet",
        ctx.qbo.fetch_report("BalanceSheet", {"start_date": start, "end_date": end}),
    )
    _report_sheet(wb, "AP Aging", ctx.qbo.fetch_report("AgedPayables", {"report_date": end}))
    _report_sheet(wb, "AR Aging", ctx.qbo.fetch_report("AgedReceivables", {"report_date": end}))
    _tie_out_sheet(wb, ctx)
    _payroll_sheet(wb, ctx)
    _owner_sheet(wb, ctx)

    # author/company metadata is the tenant's legal name (invariant 10)
    wb.properties.creator = legal_name
    wb.properties.lastModifiedBy = legal_name
    try:
        wb.properties.company = legal_name
    except AttributeError:
        pass

    target = guard.check_write(Path(out_path).expanduser())
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(target))
    return target
