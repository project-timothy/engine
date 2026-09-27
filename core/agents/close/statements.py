"""The month's financial statements: the ceremony's closing deliverable.

One workbook, three sheets — P&L for the sealed month, P&L year to date,
balance sheet at month-end — rendered straight off the accounting system's
own reports, formatted the way an owner reads statements: real numeric
cells in accounting format, section indentation, bold ruled totals, a title
block naming the company and the period. Regenerated on every render, never
hand-edited; a sealed month's statements are stable by definition.

Author/company metadata is the tenant's legal name (invariant 10); openpyxl
per the house rule; the write passes the guard. Everything here is
deterministic — the numbers are the accounting system's, verbatim.
"""

from __future__ import annotations

import calendar
from decimal import Decimal, InvalidOperation
from pathlib import Path

from ...engine.guard import WriteGuard
from .checks import CloseContext, month_bounds

_TITLE_SIZE = 14
_AMOUNT_FORMAT = "#,##0.00;(#,##0.00)"
_HEADER_FILL = "FF2F3B4C"  # slate band under the title block
_HEADER_INK = "FFFFFFFF"


def fiscal_year_start(month: str, year_start_month: int) -> str:
    """First day of the fiscal year containing ``month`` (YYYY-MM)."""
    year, mon = (int(part) for part in month.split("-"))
    if mon < year_start_month:
        year -= 1
    return f"{year:04d}-{year_start_month:02d}-01"


def _month_label(month: str) -> str:
    year, mon = (int(part) for part in month.split("-"))
    return f"{calendar.month_name[mon]} {year}"


def _ytd_label(month: str, year_start_month: int) -> str:
    year, mon = (int(part) for part in month.split("-"))
    start = fiscal_year_start(month, year_start_month)
    start_year, start_mon = int(start[:4]), int(start[5:7])
    if (start_year, start_mon) == (year, mon):
        return _month_label(month)
    if start_year == year:
        return f"{calendar.month_name[start_mon]} through {calendar.month_name[mon]} {year}"
    return (
        f"{calendar.month_name[start_mon]} {start_year} through {calendar.month_name[mon]} {year}"
    )


def _as_amount(text: str) -> Decimal | None:
    raw = text.strip().replace("$", "").replace(",", "")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def _statement_sheet(
    wb,
    *,
    sheet_title: str,
    legal_name: str,
    statement_title: str,
    period_label: str,
    report: dict,
) -> None:
    """One statement, rendered from the report's own row tree."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from ...adapters.qbo import report_rows

    ws = wb.create_sheet(sheet_title)
    columns = [
        str(c.get("ColTitle", "") or "") for c in (report.get("Columns") or {}).get("Column", [])
    ]
    width = max(2, len(columns) or 2)

    ws.append([legal_name])
    ws.cell(row=1, column=1).font = Font(bold=True, size=_TITLE_SIZE)
    ws.append([statement_title])
    ws.cell(row=2, column=1).font = Font(bold=True, size=12)
    ws.append([period_label])
    ws.append([])

    header = list(columns) if columns else ["", "Amount"]
    header[0] = ""
    ws.append(header)
    header_row = ws.max_row
    for col in range(1, width + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color=_HEADER_INK)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")

    total_rule = Border(top=Side(style="thin"))
    for row in report_rows(report):
        cells = list(row["cells"])
        if not cells:
            continue
        indent = "    " * len(row["path"])
        rendered: list[object] = [indent + cells[0]]
        for text in cells[1:]:
            amount = _as_amount(text)
            rendered.append(float(amount) if amount is not None else text)
        ws.append(rendered)
        out_row = ws.max_row
        is_total = row["kind"] == "summary"
        for col in range(2, len(rendered) + 1):
            cell = ws.cell(row=out_row, column=col)
            if isinstance(cell.value, float):
                cell.number_format = _AMOUNT_FORMAT
            if is_total:
                cell.border = total_rule
        if is_total:
            for col in range(1, len(rendered) + 1):
                existing = ws.cell(row=out_row, column=col)
                existing.font = Font(bold=True)

    label_longest = max(
        (len(str(ws.cell(row=r, column=1).value or "")) for r in range(1, ws.max_row + 1)),
        default=0,
    )
    ws.column_dimensions["A"].width = max(30, min(label_longest + 2, 60))
    for col in range(2, width + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.sheet_view.showGridLines = False


def render_statements(
    ctx: CloseContext,
    *,
    legal_name: str,
    year_start_month: int,
    out_path: str | Path,
    guard: WriteGuard,
) -> Path:
    from openpyxl import Workbook

    start, end = month_bounds(ctx.month)
    ytd_start = fiscal_year_start(ctx.month, year_start_month)
    wb = Workbook()
    wb.remove(wb.active)

    _statement_sheet(
        wb,
        sheet_title="P&L Month",
        legal_name=legal_name,
        statement_title="Profit and Loss",
        period_label=_month_label(ctx.month),
        report=ctx.qbo.fetch_report("ProfitAndLoss", {"start_date": start, "end_date": end}),
    )
    _statement_sheet(
        wb,
        sheet_title="P&L YTD",
        legal_name=legal_name,
        statement_title="Profit and Loss, Year to Date",
        period_label=_ytd_label(ctx.month, year_start_month),
        report=ctx.qbo.fetch_report("ProfitAndLoss", {"start_date": ytd_start, "end_date": end}),
    )
    year, mon = (int(part) for part in ctx.month.split("-"))
    _statement_sheet(
        wb,
        sheet_title="Balance Sheet",
        legal_name=legal_name,
        statement_title="Balance Sheet",
        period_label=f"As of {calendar.month_name[mon]} {int(end[-2:])}, {year}",
        report=ctx.qbo.fetch_report("BalanceSheet", {"start_date": start, "end_date": end}),
    )

    # author/company metadata is the tenant's legal name (invariant 10)
    wb.properties.creator = legal_name
    wb.properties.lastModifiedBy = legal_name
    try:
        wb.properties.company = legal_name
    except AttributeError:
        pass

    target = guard.check_write(Path(out_path).expanduser())
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(target))
    return target
