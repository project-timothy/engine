"""Close packet evals: refuses a blocked month, renders the evidence
workbook with the tenant's legal name in the metadata (invariant 10)."""

from __future__ import annotations

from openpyxl import load_workbook

from core.agents.close import jobs as close_jobs
from core.agents.close.checks import CloseContext
from core.agents.close.packet import render_packet
from core.agents.close.schema import CheckResult
from core.engine.guard import WriteGuard

from .test_preflight import FakeQbo, _job_ctx, _ledger, _seed_auditor

EXPECTED_SHEETS = [
    "Checklist + Sign-off",
    "P&L",
    "Balance Sheet",
    "AP Aging",
    "AR Aging",
    "Bank Tie-Out",
    "Payroll",
    "Owner Transactions",
]


def _ctx(tmp_path, ledger, *, qbo=None):
    return CloseContext(
        tenant_slug="t",
        ledger=ledger,
        month="2026-07",
        qbo=qbo or FakeQbo(),
        auditor_store=tmp_path / "aud" / "t" / "auditor.sqlite3",
        uncategorized_block_over_cents=50_000,
        statement_balance_cents=100,
    )


def _results():
    return [
        CheckResult(name="machinery", status="OK", summary="clean"),
        CheckResult(name="statement-anchor", status="WARN", summary="pending"),
    ]


def test_packet_renders_all_sheets_with_legal_name(tmp_path):
    with _ledger(tmp_path) as ledger:
        path = render_packet(
            _ctx(tmp_path, ledger),
            _results(),
            legal_name="T Corp LLC",
            out_path=tmp_path / "out" / "Close_Packet_2026-07.xlsx",
            guard=WriteGuard([]),
        )
    wb = load_workbook(path)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb.properties.creator == "T Corp LLC"
    assert wb.properties.lastModifiedBy == "T Corp LLC"
    checklist = wb["Checklist + Sign-off"]
    text = " ".join(str(c.value) for row in checklist.iter_rows() for c in row if c.value)
    assert "machinery" in text
    assert "Owner sign-off:" in text
    wb.close()


def test_tie_out_sheet_carries_both_numbers(tmp_path):
    from .test_bank_rec import _balance_sheet

    with _ledger(tmp_path) as ledger:
        qbo = FakeQbo(reports={"BalanceSheet": _balance_sheet("Checking", "1.00")})
        ctx = _ctx(tmp_path, ledger, qbo=qbo)
        ctx.bank_account = "Checking"
        path = render_packet(
            ctx,
            _results(),
            legal_name="T Corp LLC",
            out_path=tmp_path / "out" / "p.xlsx",
            guard=WriteGuard([]),
        )
    wb = load_workbook(path)
    tie = wb["Bank Tie-Out"]
    text = " ".join(str(c.value) for row in tie.iter_rows() for c in row if c.value)
    assert "$1.00" in text  # statement (100 cents) and register both render
    wb.close()


def test_packet_job_refuses_a_blocked_month(tmp_path, monkeypatch):
    monkeypatch.setattr(close_jobs, "_qbo_read_client", lambda ctx: FakeQbo())
    monkeypatch.setenv("AUDITOR_STORE_ROOT", str(tmp_path / "aud"))
    # no auditor store seeded: machinery BLOCKs
    with _ledger(tmp_path) as ledger:
        output = close_jobs._packet_run(_job_ctx(tmp_path, ledger, {"month": "2026-07"}))
    assert "NOT rendered" in output.summary
    assert "machinery" in output.summary
    assert not list((tmp_path / "reports").rglob("*.xlsx"))


def test_packet_job_renders_a_clean_month(tmp_path, monkeypatch):
    qbo = FakeQbo(
        journal_entries=[
            {"TxnDate": "2026-07-15", "DocNumber": "PayrollCo", "PrivateNote": "Regular Payroll"},
            {"TxnDate": "2026-07-31", "DocNumber": "PayrollCo", "PrivateNote": "Regular Payroll"},
        ]
    )
    monkeypatch.setattr(close_jobs, "_qbo_read_client", lambda ctx: qbo)
    monkeypatch.setenv("AUDITOR_STORE_ROOT", str(tmp_path / "aud"))
    _seed_auditor(tmp_path)
    with _ledger(tmp_path) as ledger:
        output = close_jobs._packet_run(_job_ctx(tmp_path, ledger, {"month": "2026-07"}))
    assert "rendered" in output.summary
    packet = tmp_path / "reports" / "2026-07" / "Close_Packet_2026-07.xlsx"
    assert packet.exists()
    assert output.events[0].payload["month"] == "2026-07"
