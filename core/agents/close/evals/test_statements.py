"""Statements evals: the ceremony's closing act.

Render only over a sealed month; reveal is config-gated; the send is
two-phase like the lock (park, approve, send) with the sent EVENT as the
replay guard — an approved card can never fire twice.
"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from core.agents.close import jobs as close_jobs
from core.agents.close.statements import fiscal_year_start
from core.engine.config import CloseSettings, Identity, TenantConfig
from core.engine.contracts import JobContext
from core.engine.guard import WriteGuard

from .test_lock import SealableQbo
from .test_preflight import _key, _ledger

PNL_MONTH = {
    "Columns": {"Column": [{"ColTitle": ""}, {"ColTitle": "Total"}]},
    "Rows": {
        "Row": [
            {
                "Header": {"ColData": [{"value": "Income"}]},
                "Rows": {"Row": [{"ColData": [{"value": "Services"}, {"value": "255300.00"}]}]},
                "Summary": {"ColData": [{"value": "Total Income"}, {"value": "255300.00"}]},
            },
            {"ColData": [{"value": "Net Income"}, {"value": "139353.18"}]},
        ]
    },
}

BALANCE_SHEET = {
    "Columns": {"Column": [{"ColTitle": ""}, {"ColTitle": "Total"}]},
    "Rows": {
        "Row": [
            {
                "Header": {"ColData": [{"value": "Assets"}]},
                "Rows": {"Row": [{"ColData": [{"value": "Checking"}, {"value": "258559.95"}]}]},
                "Summary": {"ColData": [{"value": "Total Assets"}, {"value": "258559.95"}]},
            }
        ]
    },
}

REPORTS = {"ProfitAndLoss": PNL_MONTH, "BalanceSheet": BALANCE_SHEET}


class FakeMailer:
    def __init__(self):
        self.sent: list[dict] = []

    def send_mail(self, *, subject, body, to, attachments=()):
        self.sent.append(
            {"subject": subject, "body": body, "to": list(to), "attachments": list(attachments)}
        )


def _job_ctx(tmp_path, ledger, params=None, *, recipients=("reviewer@example.com",), reveal=False):
    tenant = TenantConfig(
        identity=Identity(legal_name="T Corp", slug="t", timezone="UTC"),
        close=CloseSettings(
            report_dir=str(tmp_path / "reports"),
            statements_recipients=list(recipients),
            reveal_in_finder=reveal,
        ),
    )
    return JobContext(
        tenant=tenant,
        tenant_slug="t",
        ledger=ledger,
        agent="close",
        job="statements",
        params=params or {"month": "2026-07"},
        guard=WriteGuard([]),
    )


def _card(ledger, *, month="2026-07", status="pending"):
    run_id = ledger.conn.execute(
        "INSERT INTO runs (idempotency_key, tenant, agent, job, status, result_json, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (_key("run"), "t", "close", "statements", "ok", "{}", "2026-08-03T06:00:00+00:00"),
    ).lastrowid
    cursor = ledger.conn.execute(
        "INSERT INTO approval_queue (idempotency_key, run_id, tenant, agent, action_type, "
        "params_json, status, created_at) VALUES (?,?,?,?,?,?,?,?)",
        (
            _key("appr"),
            run_id,
            "t",
            "close",
            close_jobs.STATEMENTS_ACTION,
            json.dumps({"month": month}),
            status,
            "2026-08-03T06:00:00+00:00",
        ),
    )
    ledger.conn.commit()
    return int(cursor.lastrowid)


def _sent_event(ledger, *, month="2026-07", card_id):
    run_id = ledger.conn.execute(
        "INSERT INTO runs (idempotency_key, tenant, agent, job, status, result_json, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (_key("run"), "t", "close", "statements", "ok", "{}", "2026-08-03T06:00:00+00:00"),
    ).lastrowid
    ledger.conn.execute(
        "INSERT INTO events (idempotency_key, run_id, tenant, agent, event_type, payload_json, "
        "created_at) VALUES (?,?,?,?,?,?,?)",
        (
            _key("evt"),
            run_id,
            "t",
            "close",
            close_jobs.STATEMENTS_SENT_EVENT,
            json.dumps({"month": month, "card_id": card_id}),
            "2026-08-03T06:00:00+00:00",
        ),
    )
    ledger.conn.commit()


def _run(tmp_path, ledger, monkeypatch, *, sealed=True, ctx=None, mailer=None, mailer_factory=None):
    """``mailer`` replaces the client the factory hands back; ``mailer_factory``
    replaces the factory itself (for the pre-transport failure shapes: the
    factory is where the token is acquired)."""
    qbo = SealableQbo(
        book_close_date="2026-07-31" if sealed else "2026-06-30",
        reports=REPORTS,
    )
    mailer = mailer if mailer is not None else FakeMailer()
    monkeypatch.setattr(close_jobs, "_qbo_read_client", lambda c: qbo)
    monkeypatch.setattr(close_jobs, "_graph_send_client", mailer_factory or (lambda c: mailer))
    monkeypatch.setenv("AUDITOR_STORE_ROOT", str(tmp_path / "aud"))
    ctx = ctx or _job_ctx(tmp_path, ledger)
    return close_jobs._statements_run(ctx), mailer


def test_unsealed_month_refuses_to_render(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        output, mailer = _run(tmp_path, ledger, monkeypatch, sealed=False)
    assert "not sealed" in output.summary
    assert not (tmp_path / "reports" / "2026-07").exists()
    assert mailer.sent == []


def test_sealed_month_renders_statements_and_parks_card(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert output.status == "needs_approval"
    card = output.approvals[0]
    assert card.action_type == close_jobs.STATEMENTS_ACTION
    assert card.params["month"] == "2026-07"
    assert mailer.sent == []  # parking never sends
    path = tmp_path / "reports" / "2026-07" / "Financial_Statements_2026-07.xlsx"
    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["P&L Month", "P&L YTD", "Balance Sheet"]
    ws = wb["P&L Month"]
    assert ws["A1"].value == "T Corp"
    amounts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, float)]
    assert 139353.18 in amounts  # real numeric cells, not text
    assert wb.properties.creator == "T Corp"


def test_pending_card_renders_but_waits(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="pending")
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert f"awaiting approval (card #{card_id})" in output.summary
    assert mailer.sent == []


def test_approved_card_sends_with_attachment_and_records_event(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert "SENT" in output.summary
    assert len(mailer.sent) == 1
    message = mailer.sent[0]
    assert message["to"] == ["reviewer@example.com"]
    assert "July 2026" in message["subject"]
    name, mime, data = message["attachments"][0]
    assert name == "Financial_Statements_2026-07.xlsx"
    assert "spreadsheetml" in mime
    assert len(data) > 1000  # the real workbook bytes ride along
    sent = [e for e in output.events if e.event_type == close_jobs.STATEMENTS_SENT_EVENT]
    assert sent[0].payload["card_id"] == card_id


def test_sent_event_blocks_a_second_send(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")
        _sent_event(ledger, card_id=card_id)
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert "already sent" in output.summary
    assert mailer.sent == []


def test_fresh_card_after_send_allows_a_fresh_send(tmp_path, monkeypatch):
    # A new approved card (new id) is a new ask: the old sent event does not
    # block it. This is how the owner re-sends corrected statements.
    with _ledger(tmp_path) as ledger:
        old_card = _card(ledger, status="approved")
        _sent_event(ledger, card_id=old_card)
        ledger.conn.execute("UPDATE approval_queue SET status='rejected' WHERE id=?", (old_card,))
        ledger.conn.commit()
        _card(ledger, status="approved")
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert "SENT" in output.summary
    assert len(mailer.sent) == 1


def test_shadow_never_renders_or_sends(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        _card(ledger, status="approved")
        ctx = _job_ctx(tmp_path, ledger)
        ctx.shadow = True
        output, mailer = _run(tmp_path, ledger, monkeypatch, ctx=ctx)
    assert "would render" in output.summary
    assert not (tmp_path / "reports" / "2026-07").exists()
    assert mailer.sent == []


def test_no_recipients_renders_and_skips_the_send_leg(tmp_path, monkeypatch):
    with _ledger(tmp_path) as ledger:
        ctx = _job_ctx(tmp_path, ledger, recipients=())
        output, mailer = _run(tmp_path, ledger, monkeypatch, ctx=ctx)
    assert "send skipped" in output.summary
    assert output.status == "ok"
    assert not output.approvals
    assert (tmp_path / "reports" / "2026-07" / "Financial_Statements_2026-07.xlsx").exists()


def test_reveal_is_config_gated(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(close_jobs, "_reveal_in_finder", lambda p: calls.append(p) or True)
    with _ledger(tmp_path) as ledger:
        _run(tmp_path, ledger, monkeypatch)  # default ctx: reveal off
    assert calls == []
    with _ledger(tmp_path / "b") as ledger:
        ctx = _job_ctx(tmp_path / "b", ledger, reveal=True)
        output, _ = _run(tmp_path / "b", ledger, monkeypatch, ctx=ctx)
    assert len(calls) == 1
    assert "revealed in Finder" in output.actions


def test_key_changes_with_card_state(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDITOR_STORE_ROOT", str(tmp_path / "aud"))
    with _ledger(tmp_path) as ledger:
        ctx = _job_ctx(tmp_path, ledger)
        key_none = close_jobs._statements_key(ctx)
        _card(ledger, status="pending")
        key_pending = close_jobs._statements_key(ctx)
    assert key_none != key_pending


def test_fiscal_year_start_handles_offset_years():
    assert fiscal_year_start("2026-07", 1) == "2026-01-01"
    assert fiscal_year_start("2026-07", 10) == "2025-10-01"
    assert fiscal_year_start("2026-11", 10) == "2026-10-01"


# ---- issue #135: the send is at-most-once ------------------------------------


def _card_params(ledger, card_id):
    row = ledger.conn.execute(
        "SELECT params_json FROM approval_queue WHERE id = ?", (card_id,)
    ).fetchone()
    return json.loads(row["params_json"])


def test_2026_08_20_interrupted_send_never_auto_resends(tmp_path, monkeypatch):
    """Issue #135: the Graph accept and the sent event are two steps; a
    crash between them must never become a second email. send_started is
    stamped on the card BEFORE the transport call, so a retry sees an
    attempt with no recorded outcome, refuses to resend, and parks a fresh
    card (fresh card = deliberate re-send, the standing doctrine)."""
    import pytest as _pytest

    class ExplodingMailer:
        calls = 0

        def send_mail(self, **kwargs):
            type(self).calls += 1
            raise RuntimeError("transport died after the message may have gone out")

    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")
        with _pytest.raises(RuntimeError):
            _run(tmp_path, ledger, monkeypatch, mailer=ExplodingMailer())
        assert ExplodingMailer.calls == 1
        assert _card_params(ledger, card_id).get("send_started")  # stamped pre-send

        output, mailer = _run(tmp_path, ledger, monkeypatch)

        assert mailer.sent == []  # never an automatic second send
        assert any(a.code == "close.statements_send_unconfirmed" for a in output.anomalies)
        assert output.status == "needs_approval"
        fresh = [a for a in output.approvals if a.action_type == close_jobs.STATEMENTS_ACTION]
        assert len(fresh) == 1  # the deliberate re-send path


def test_completed_send_with_marker_still_reads_already_sent(tmp_path, monkeypatch):
    """The marker plus the sent event is the NORMAL completed state; it must
    read as already-sent, never as unconfirmed. (The sent event is seeded by
    hand because this harness calls the job directly and the runner is what
    persists a JobOutput's events.)"""
    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")
        output1, mailer1 = _run(tmp_path, ledger, monkeypatch)
        assert len(mailer1.sent) == 1  # the send itself still works
        assert _card_params(ledger, card_id).get("send_started")  # marker stamped
        _sent_event(ledger, card_id=card_id)

        output2, mailer2 = _run(tmp_path, ledger, monkeypatch)

        assert mailer2.sent == []
        assert "already sent" in output2.summary
        assert output2.anomalies == []


def test_rejected_send_card_reparks_on_a_fresh_key(tmp_path, monkeypatch):
    """Incident 2026-09-01 (issue #161): the documented re-send path is
    reject-the-old-card-approve-a-new-one, but the fresh park rode the same
    stable key ("statements:{month}") and the runner's #143 dedup swallowed
    it against the rejected row. The re-ask must name the card it supersedes."""
    with _ledger(tmp_path) as ledger:
        rejected_id = _card(ledger, status="rejected")
        output, mailer = _run(tmp_path, ledger, monkeypatch)
    assert output.status == "needs_approval"
    assert output.approvals[0].key == f"statements:2026-07:reask-{rejected_id}"
    assert mailer.sent == []


# ---- honesty audit 2026-09-03: claims match outcomes ------------------------


def test_reveal_claims_nothing_when_open_fails(monkeypatch):
    """F14: ``open -R`` exits non-zero in a headless launchd context or on a
    missing path; the helper used to return True regardless, and the run
    claimed "revealed in Finder". The claim follows the return code."""
    import subprocess
    import sys

    monkeypatch.setattr(sys, "platform", "darwin")
    seen = []

    def fake_run(argv, **kwargs):
        seen.append(argv)
        return subprocess.CompletedProcess(argv, returncode=1)

    monkeypatch.setattr(subprocess, "run", fake_run)
    assert close_jobs._reveal_in_finder("/nowhere/Financial_Statements_2026-07.xlsx") is False
    assert seen and seen[0][:2] == ["/usr/bin/open", "-R"]

    monkeypatch.setattr(
        subprocess, "run", lambda argv, **kw: subprocess.CompletedProcess(argv, returncode=0)
    )
    assert close_jobs._reveal_in_finder("/somewhere/Financial_Statements_2026-07.xlsx") is True


def test_a_linux_host_never_spawns_the_macos_open(monkeypatch):
    """Row 7.20: ``open -R`` is the one macOS binary on an engine job's path.
    Off Darwin the helper answers False without spawning anything, so the
    container's close renders the workbook and says nothing about Finder."""
    import subprocess
    import sys

    monkeypatch.setattr(sys, "platform", "linux")

    def never(*args, **kwargs):  # pragma: no cover - the point is that it is not called
        raise AssertionError(f"a subprocess was spawned off Darwin: {args}")

    monkeypatch.setattr(subprocess, "run", never)
    assert close_jobs._reveal_in_finder("/srv/engine/Financial_Statements_2026-07.xlsx") is False


def test_auth_failure_before_the_stamp_leaves_no_stamp_and_the_next_run_sends(
    tmp_path, monkeypatch
):
    """F11: the mail client's token is acquired lazily inside the transport
    call, so an absent/expired MSAL cache raised AFTER send_started was
    stamped; the next run then said "a send started but its outcome was
    never recorded", sent the owner to Sent Items for a mail never
    attempted, and parked a resend card. The client and its token now come
    BEFORE the stamp: an auth failure stamps nothing, reports itself as a
    send not attempted, and the next run sends normally."""
    from core.adapters.graph_mail import GraphAuthError

    def no_token(ctx):
        raise GraphAuthError("no MSAL token cache in keychain (svc); run device-code auth")

    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")
        output1, _ = _run(tmp_path, ledger, monkeypatch, mailer_factory=no_token)
        assert output1.status == "error"
        assert "NOT attempted" in output1.summary
        assert any(a.code == "close.statements_send_not_attempted" for a in output1.anomalies)
        assert "send_started" not in _card_params(ledger, card_id)  # nothing stamped
        assert not output1.approvals  # no resend card: there was nothing to resend

        output2, mailer = _run(tmp_path, ledger, monkeypatch)

        assert "SENT" in output2.summary
        assert len(mailer.sent) == 1
        assert output2.anomalies == []
        assert _card_params(ledger, card_id).get("send_started")


def test_client_comes_before_the_stamp_and_the_stamp_before_the_transport(tmp_path, monkeypatch):
    """The at-most-once order, pinned end to end: build the client (token
    in hand) -> stamp send_started -> POST. The client factory must see no
    stamp; the transport must see one (issue #135 doctrine unchanged)."""
    order = []

    with _ledger(tmp_path) as ledger:
        card_id = _card(ledger, status="approved")

        class OrderMailer:
            def send_mail(self, **kwargs):
                order.append(("send", bool(_card_params(ledger, card_id).get("send_started"))))

        def factory(ctx):
            order.append(("client", bool(_card_params(ledger, card_id).get("send_started"))))
            return OrderMailer()

        output, _ = _run(tmp_path, ledger, monkeypatch, mailer_factory=factory)
    assert "SENT" in output.summary
    assert order == [("client", False), ("send", True)]
