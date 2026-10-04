"""Workbook view: the human-readable AP ledger, generated from the engine ledger.

A delivery view (invariant 1: the ledger is truth; this is a regenerated
photograph of it, never hand-edited). The column layout is tenant config so the
sheet matches the tenant's legacy workbook column-for-column while core stays
tenant-agnostic. Rows are colored by status group, the operator's at-a-glance
unpaid / scheduled / paid (2026-06-26):

- blank  = payable group (Received / Approved / Outstanding): nothing has happened
- red    = committed group (Scheduled): payment exists, not yet cleared
- green  = Paid: cleared the bank
- grey   = Void / Cancelled: settled but never cleared

openpyxl per the house rule; the author metadata is the tenant's legal name
(invariant 10). The output path passes the write guard, so a workbook_path
inside a protected production surface is refused.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ...engine.config import WorkbookColumn
from ...engine.guard import WriteGuard
from ...ledger import Ledger
from .status import is_committed, is_payable_eligible, is_settled

# Bump whenever the rendered output changes (layout, colors, styling). The
# workbook job folds this into its idempotency key, so a renderer change forces
# a regenerate even when the ledger data is unchanged (without it, a code change
# silently leaves a stale view on disk).
VIEW_VERSION = "5"  # 5: Vendor Summary sheet (issue #166)

# Standard Excel status fills (ARGB), instantly recognizable.
_RED = "FFFFC7CE"  # scheduled (committed)
_GREEN = "FFC6EFCE"  # paid (cleared the bank)
_GREY = "FFD9D9D9"  # void / cancelled (settled, never cleared)


def _status_fill(status: str) -> str | None:
    """Fill color for a status, by group. None (blank) = nothing has happened."""
    if is_payable_eligible(status):
        return None
    if is_committed(status):
        return _RED
    if status == "Paid":
        return _GREEN
    if is_settled(status):  # Void / Cancelled: settled but not a bank clear
        return _GREY
    return None


# Buckets in display order, top to bottom: what needs action first, done last.
# Paid is split from void/cancelled so the cash-flow subtotal counts real cash.
_BUCKETS = [
    ("pending", "Pending"),
    ("scheduled", "Scheduled"),
    ("paid", "Paid"),
    ("closed", "Void / Cancelled"),
]


def _bucket(status: str) -> str:
    if is_payable_eligible(status):
        return "pending"
    if is_committed(status):
        return "scheduled"
    if status == "Paid":
        return "paid"
    return "closed"  # void / cancelled: no cash meaning


def _cell_value(row: dict, field: str) -> Any:
    if not field:
        return ""  # a column with no engine field (e.g. Payment Date) renders blank
    if field == "amount_cents":
        cents = row.get("amount_cents")
        return f"${cents / 100:,.2f}" if cents is not None else ""
    return row.get(field, "") or ""


def render_workbook(
    rows: list[dict],
    columns: list[WorkbookColumn],
    *,
    legal_name: str,
    unprocessed: list[dict] | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "AP Ledger"

    ws.append([c.label for c in columns])

    amount_idx = next(
        (i for i, c in enumerate(columns, start=1) if c.field == "amount_cents"), None
    )
    status_idx = next((i for i, c in enumerate(columns, start=1) if c.field == "status"), 2)
    grouped: dict[str, list[dict]] = {b: [] for b, _ in _BUCKETS}
    for row in rows:
        grouped[_bucket(str(row.get("status", "")))].append(row)

    for bucket, label in _BUCKETS:
        group = sorted(
            grouped[bucket],
            key=lambda r: (str(r.get("vendor", "")), str(r.get("invoice_number", ""))),
        )
        if not group:
            continue
        for row in group:
            ws.append([_cell_value(row, c.field) for c in columns])
            color = _status_fill(str(row.get("status", "")))
            if color:
                fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                for cell in ws[ws.max_row]:
                    cell.fill = fill
        # cash-flow subtotal under each cash-meaningful bucket (not void/cancelled)
        if bucket != "closed" and amount_idx is not None:
            total = sum((r.get("amount_cents") or 0) for r in group)
            line = [""] * len(columns)
            line[0] = f"{label} total ({len(group)})"
            line[amount_idx - 1] = f"${total / 100:,.2f}"
            ws.append(line)
            subtotal_fill = PatternFill("solid", fgColor="FFF2F2F2")
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True, italic=True)
                cell.fill = subtotal_fill

    if unprocessed:
        # The "needs identification" pile: files the engine could not turn into
        # an invoice. The owner dismisses (junk) or identifies (a missed invoice)
        # each; a resolved one never reappears (see unprocessed.py).
        divider = [""] * len(columns)
        divider[0] = f"Unprocessed / needs identification ({len(unprocessed)})"
        ws.append(divider)
        header_fill = PatternFill("solid", fgColor="FFFFE699")  # amber: needs your eyes
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for item in unprocessed:
            line = [""] * len(columns)
            line[0] = item.get("file", "")
            line[status_idx - 1] = item.get("reason", "needs identification")
            ws.append(line)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="FFFFF2CC")  # light amber

    _style_worksheet(ws, columns)
    _render_vendor_summary(wb, rows)
    wb.active = 0  # the daily view opens on the AP Ledger sheet, where it always did

    # author/company metadata is the tenant's legal name (invariant 10)
    wb.properties.creator = legal_name
    wb.properties.lastModifiedBy = legal_name
    return wb


_SUMMARY_HEADERS = [
    "Vendor",
    "# Invoices",
    "Total Paid",
    "Outstanding (payable)",
    "Committed (scheduled)",
]


def _render_vendor_summary(wb, rows: list[dict]) -> None:
    """Second sheet: per-vendor rollup of the same rows the AP Ledger sheet shows.

    Regenerated with the workbook, so it can never be older than the ledger
    sheet beside it (a hand-copied rollup elsewhere rots; issue #166). Grouping
    is by vendor spelling as recorded; taxpayer/1099 grouping stays the
    auditor's year-end appendix. Void/cancelled rows carry no cash meaning and
    are excluded from counts and sums. Amounts are numeric cells so a reader
    can sum the columns.
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    agg: dict[str, list[int]] = {}  # vendor -> [count, paid, outstanding, committed] (cents)
    for row in rows:
        status = str(row.get("status", ""))
        cents = row.get("amount_cents") or 0
        if is_settled(status) and status != "Paid":
            continue  # void / cancelled: no cash meaning
        entry = agg.setdefault(str(row.get("vendor", "")), [0, 0, 0, 0])
        entry[0] += 1
        if status == "Paid":
            entry[1] += cents
        elif is_payable_eligible(status):
            entry[2] += cents
        elif is_committed(status):
            entry[3] += cents

    ws = wb.create_sheet("Vendor Summary")
    ws.append(_SUMMARY_HEADERS)
    ranked = sorted(agg.items(), key=lambda kv: (-(kv[1][1] + kv[1][2] + kv[1][3]), kv[0]))
    for vendor, (count, paid, outstanding, committed) in ranked:
        ws.append([vendor, count, paid / 100, outstanding / 100, committed / 100])
    ws.append(
        [
            "TOTAL",
            sum(entry[0] for entry in agg.values()),
            sum(entry[1] for entry in agg.values()) / 100,
            sum(entry[2] for entry in agg.values()) / 100,
            sum(entry[3] for entry in agg.values()) / 100,
        ]
    )

    currency = "$#,##0.00"
    for r in range(2, ws.max_row + 1):
        for idx in (3, 4, 5):
            ws.cell(row=r, column=idx).number_format = currency
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFD9E1F2")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="FF8EA9DB"))
    for idx, header in enumerate(_SUMMARY_HEADERS, start=1):
        longest = max(
            (len(str(ws.cell(row=r, column=idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[get_column_letter(idx)].width = max(len(header), longest) + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _style_worksheet(ws, columns: list[WorkbookColumn]) -> None:
    """Make the sheet readable, and keep it readable as it grows: content-fit
    column widths (capped, with wrapping past the cap so nothing overflows), a
    frozen and filterable header row, row separators, and right-aligned amounts.
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    width_cap, width_min = 42, 9
    wrap = Alignment(wrap_text=True, vertical="top")
    wrap_right = Alignment(wrap_text=True, vertical="top", horizontal="right")
    row_rule = Border(bottom=Side(style="thin", color="FFD9D9D9"))
    amount_cols = {i for i, c in enumerate(columns, start=1) if c.field == "amount_cents"}

    for idx in range(1, len(columns) + 1):
        longest = max(
            (len(str(ws.cell(row=r, column=idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[get_column_letter(idx)].width = max(
            width_min, min(longest + 2, width_cap)
        )

    for r in range(2, ws.max_row + 1):
        for idx in range(1, len(columns) + 1):
            cell = ws.cell(row=r, column=idx)
            cell.alignment = wrap_right if idx in amount_cols else wrap
            cell.border = row_rule

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFD9E1F2")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="FF8EA9DB"))
    ws.freeze_panes = "A2"  # header stays visible as the ledger grows
    ws.auto_filter.ref = ws.dimensions  # sortable / filterable


def ledger_rows(ledger: Ledger, tenant: str) -> list[dict]:
    """Every AP invoice the engine holds for the tenant, for the view."""
    rows = ledger.conn.execute(
        "SELECT * FROM ap_invoices WHERE tenant = ? ORDER BY id", (tenant,)
    ).fetchall()
    return [dict(r) for r in rows]


def run_workbook(
    *,
    tenant_slug: str,
    ledger: Ledger,
    columns: list[WorkbookColumn],
    legal_name: str,
    guard: WriteGuard,
    out_path: str | Path,
) -> Path:
    """Render the workbook view and write it inside an allowed (non-protected) path."""
    from .unprocessed import unresolved_unprocessed

    target = guard.check_write(out_path)  # refuses a protected production surface
    wb = render_workbook(
        ledger_rows(ledger, tenant_slug),
        columns,
        legal_name=legal_name,
        unprocessed=unresolved_unprocessed(ledger, tenant_slug),
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(target))
    return target
